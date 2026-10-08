import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_human_cti_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Professional Slate / Obsidian Palette (Non-AI, Mandiant/SANS style)
    COLOR_BG = RGBColor(11, 15, 23)           # Deep Obsidian #0B0F17
    COLOR_SURFACE = RGBColor(18, 24, 38)      # Card surface #121826
    COLOR_SURFACE_ALT = RGBColor(24, 32, 50)  # Lighter surface #182032
    COLOR_BORDER = RGBColor(40, 53, 80)       # Crisp border #283550
    COLOR_BORDER_LIGHT = RGBColor(60, 78, 115)
    
    # Typography colors
    COLOR_TEXT_PRIMARY = RGBColor(241, 245, 249)    # Slate 100
    COLOR_TEXT_SECONDARY = RGBColor(148, 163, 184)  # Slate 400
    COLOR_TEXT_MUTED = RGBColor(100, 116, 139)      # Slate 500
    
    # Precise Tactical Accents (Muted & Executive)
    ACCENT_CYAN = RGBColor(56, 189, 248)      # Sky 400
    ACCENT_BLUE = RGBColor(96, 165, 250)      # Blue 400
    ACCENT_AMBER = RGBColor(251, 191, 36)     # Amber 400
    ACCENT_RED = RGBColor(248, 113, 113)      # Red 400
    ACCENT_GREEN = RGBColor(52, 211, 153)     # Emerald 400
    ACCENT_PURPLE = RGBColor(192, 132, 252)   # Purple 400

    def apply_canvas(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()

        # Top subtle header baseline
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.3), Inches(11.733), Pt(1))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_BORDER
        line.line.fill.background()
        return bg

    def render_header(slide, section_tag, title_text, subtitle_text=""):
        # Section breadcrumb / TLP tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(8.0), Inches(0.3))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = section_tag.upper()
        p_tag.font.name = "Segoe UI"
        p_tag.font.size = Pt(9)
        p_tag.font.bold = True
        p_tag.font.color.rgb = ACCENT_CYAN

        # Right-aligned document status
        status_box = slide.shapes.add_textbox(Inches(9.0), Inches(0.35), Inches(3.5), Inches(0.3))
        tf_stat = status_box.text_frame
        p_stat = tf_stat.paragraphs[0]
        p_stat.text = "TLP:CLEAR  //  THREAT INTEL BRIEFING"
        p_stat.font.name = "Segoe UI"
        p_stat.font.size = Pt(8.5)
        p_stat.font.bold = True
        p_stat.font.color.rgb = COLOR_TEXT_MUTED
        p_stat.alignment = PP_ALIGN.RIGHT

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.733), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Segoe UI"
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_PRIMARY

    def draw_box(slide, left, top, width, height, bg_color=COLOR_SURFACE, border_color=COLOR_BORDER, border_width=1):
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        box.fill.solid()
        box.fill.fore_color.rgb = bg_color
        if border_color:
            box.line.color.rgb = border_color
            box.line.width = Pt(border_width)
        else:
            box.line.fill.background()
        return box

    # =========================================================================
    # SLIDE 1: FORMAL BRIEFING COVER (HUMAN EXECUTIVE DESIGN)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_BG
    bg1.line.fill.background()

    # Left accent architectural stripe
    stripe = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Pt(4), Inches(4.5))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = ACCENT_CYAN
    stripe.line.fill.background()

    # Classification & series chip
    c_chip = s1.shapes.add_textbox(Inches(1.1), Inches(1.4), Inches(10.0), Inches(0.35))
    tf = c_chip.text_frame
    p = tf.paragraphs[0]
    p.text = "SECURITY BRIEFING SERIES  //  FOUNDATIONAL MODULE"
    p.font.name = "Segoe UI"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    # Main Title
    c_title = s1.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(11.0), Inches(2.2))
    tf = c_title.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Cyber Threat Intelligence (CTI)"
    p.font.name = "Segoe UI"
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_PRIMARY

    p2 = tf.add_paragraph()
    p2.text = "Prinsip Operasional, Model Analisis, dan Ekosistem Pertahanan Berbasis Ancaman"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(17)
    p2.font.color.rgb = COLOR_TEXT_SECONDARY
    p2.space_before = Pt(8)

    p3 = tf.add_paragraph()
    p3.text = "Sebuah panduan komprehensif bagi praktisi keamanan siber, analis SOC, dan pengambil keputusan teknis."
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(12)
    p3.font.color.rgb = COLOR_TEXT_MUTED
    p3.space_before = Pt(16)

    # Metadata card at bottom
    draw_box(s1, Inches(1.1), Inches(4.8), Inches(11.2), Inches(1.5), COLOR_SURFACE, COLOR_BORDER)
    meta_box = s1.shapes.add_textbox(Inches(1.3), Inches(4.9), Inches(10.8), Inches(1.3))
    tf_m = meta_box.text_frame
    tf_m.word_wrap = True
    
    pm1 = tf_m.paragraphs[0]
    pm1.text = "CAKUPAN PEMBAHASAN:"
    pm1.font.name = "Segoe UI"
    pm1.font.size = Pt(9.5)
    pm1.font.bold = True
    pm1.font.color.rgb = ACCENT_AMBER

    pm2 = tf_m.add_paragraph()
    pm2.text = "01. Kesenjangan Pertahanan Reaktif  |  02. Taksonomi Data vs Intelijen  |  03. 4 Tingkatan Konsumen CTI\n04. Siklus Intelijen 6 Tahap  |  05. Diamond Model & Kill Chain  |  06. Pyramid of Pain & Biaya Lawan\n07. Active Deception & Honeytokens  |  08. Atribusi Teknis & STIX/TAXII  |  09. Otomasi Analisis Modern"
    pm2.font.name = "Segoe UI"
    pm2.font.size = Pt(10)
    pm2.font.color.rgb = COLOR_TEXT_SECONDARY
    pm2.space_before = Pt(4)

    # Footer classification
    foot = s1.shapes.add_textbox(Inches(1.1), Inches(6.5), Inches(11.2), Inches(0.4))
    tf_f = foot.text_frame
    pf = tf_f.paragraphs[0]
    pf.text = "DISTRIBUSI: TLP:CLEAR  •  TARGET AUDIENS: SOC L1-L3, THREAT HUNTER, SECURITY LEADERSHIP"
    pf.font.name = "Segoe UI"
    pf.font.size = Pt(9)
    pf.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 2: PERMASALAHAN PERTAHANAN REAKTIF (THE ASYMMETRIC PROBLEM)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_canvas(s2)
    render_header(s2, "Konteks Industri & Masalah Lapangan", "Kelemahan Model Pertahanan Reaktif: Mengapa Perimeter Saja Tidak Cukup?")

    # Left: The Defensive Dilemma (3 metrics/facts)
    draw_box(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3), COLOR_SURFACE, COLOR_BORDER)
    tb = s2.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.9))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "FAKTA LAPANGAN: ASIMETRI KEAMANAN SIBER"
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED

    points_left = [
        ("Median Dwell Time Tinggi", "Penyerang rata-rata berada di dalam jaringan korban selama berhari-hari hingga berminggu-minggu sebelum terdeteksi secara tidak sengaja."),
        ("Alert Fatigue Kronis", "SOC enterprise menerima 5.000 hingga 15.000 peringatan per hari. Lebih dari 60% alarm diabaikan karena rasio false-positive yang masif."),
        ("Kerugian Tak Terelakkan", "Deteksi reaktif (menunggu alert SIEM berbunyi) berarti malware telah berhasil dieksekusi atau data penting telah mulai dienkripsi.")
    ]
    for title, desc in points_left:
        pt = tf.add_paragraph()
        pt.text = title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(12)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_TEXT_PRIMARY
        pt.space_before = Pt(14)

        pd = tf.add_paragraph()
        pd.text = desc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = COLOR_TEXT_SECONDARY
        pd.space_before = Pt(3)

    # Right: The CTI Shift (3 strategic changes)
    draw_box(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3), COLOR_SURFACE, COLOR_BORDER)
    tb_r = s2.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.9))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "PERUBAHAN PARADIGMA: THREAT-INFORMED DEFENSE"
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    points_right = [
        ("Fokus Keluar (External Visibility)", "Bukan hanya melihat log server sendiri, tetapi memetakan infrastruktur penyerang, kampanye malware aktif, dan sektor bisnis yang ditargetkan."),
        ("Prioritas Berbasis Relevansi", "Alih-alih memblokir jutaan IP acak, fokuskan energi pada TTP (Taktik, Teknik, Prosedur) dari aktor ancaman yang memang menargetkan industri kita."),
        ("Pencegahan di Tahap Awal", "Memutus serangan pada fase pengintaian (Reconnaissance) atau pengiriman (Delivery), jauh sebelum penyerang mencapai tujuannya.")
    ]
    for title, desc in points_right:
        pt = tf_r.add_paragraph()
        pt.text = title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(12)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_TEXT_PRIMARY
        pt.space_before = Pt(14)

        pd = tf_r.add_paragraph()
        pd.text = desc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(10.5)
        pd.font.color.rgb = COLOR_TEXT_SECONDARY
        pd.space_before = Pt(3)

    # =========================================================================
    # SLIDE 3: DATA VS INFORMASI VS INTELIJEN (TABLE OF TRANSFORMATION)
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_canvas(s3)
    render_header(s3, "Taksonomi & Definisi", "Data Mentah vs Informasi vs Intelijen Ancaman (CTI)")

    # 3 Horizontal Comparison Rows
    tiers_data = [
        ("TINGKAT 1: DATA MENTAH", "Catatan fakta atomik tanpa konteks, belum diproses, belum dapat dijadikan dasar tindakan.", 
         "Contoh Kasus:\nLog firewall mencatat adanya koneksi masuk dari 198.51.100.25 pada port 443 pada jam 03:12:10 UTC.",
         "Nilai Operasional: Rendah (Hanya mencatat kejadian masa lalu)", COLOR_TEXT_MUTED, Inches(1.6)),
        ("TINGKAT 2: INFORMASI", "Data yang telah diagregasi, dikelompokkan, dan diperkaya dengan metadata struktural.", 
         "Contoh Kasus:\nIP 198.51.100.25 terdeteksi melakukan SYN scan ke 85 port dalam 3 detik, lalu mencoba autentikasi SSH sebanyak 12 kali.",
         "Nilai Operasional: Menengah (Menjelaskan aktivitas teknis)", ACCENT_BLUE, Inches(3.3)),
        ("TINGKAT 3: INTELIJEN (CTI)", "Informasi yang telah diverifikasi, dianalisis motif dan pelakunya, serta menuntun keputusan nyata organisasi.", 
         "Contoh Kasus:\nIP tersebut adalah exit-proxy yang digunakan kelompok APT29 untuk kampanye CVE-2024-xxxx pada instansi finansial.\nTindakan Taktis: Terapkan mitigasi virtual patching, putus sesi aktif, dan monitor pergerakan lateral internal.",
         "Nilai Operasional: Kritis & Actionable (Mencegah kerugian nyata)", ACCENT_GREEN, Inches(5.0))
    ]

    for tag, desc, example, value_label, border_c, top_p in tiers_data:
        draw_box(s3, Inches(0.8), top_p, Inches(11.733), Inches(1.5), COLOR_SURFACE, border_c)
        box = s3.shapes.add_textbox(Inches(1.0), top_p + Inches(0.1), Inches(11.3), Inches(1.3))
        tf = box.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = f"{tag}   |   {value_label}"
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = border_c

        p1 = tf.add_paragraph()
        p1.text = desc
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(10.5)
        p1.font.color.rgb = COLOR_TEXT_PRIMARY
        p1.space_before = Pt(2)

        p2 = tf.add_paragraph()
        p2.text = example
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = COLOR_TEXT_SECONDARY
        p2.space_before = Pt(3)

    # =========================================================================
    # SLIDE 4: 4 TINGKATAN KONSUMEN CTI (THE 4 TIERS)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_canvas(s4)
    render_header(s4, "Arsitektur Pengambilan Keputusan", "Empat Tingkatan CTI: Memetakan Audiens, Cakupan, dan Format Deliverable")

    tiers_4 = [
        ("STRATEGIC CTI", "Pengambil Keputusan Eksekutif", "CISO, Direksi, Dewan Komisaris",
         "Tren risiko jangka panjang, motif geopolitik ancaman, estimasi kerugian finansial, dan prioritas alokasi anggaran belanja keamanan siber.",
         "Format Output:\nExecutive Summary 2 Halaman, Presentasi Risiko Dewan Direksi, Threat Landscape Report tahunan.",
         ACCENT_PURPLE, Inches(0.8)),
        ("OPERATIONAL CTI", "Manajemen Operasi Keamanan", "SOC Manager, Lead Threat Hunter, DFIR",
         "Profil spesifik kelompok penyerang (Threat Actor Dossier), kampanye penyerangan aktif, kapabilitas teknis lawan, dan jendela waktu serangan.",
         "Format Output:\nThreat Actor Dossier, Analisis Kampanye, Profil Taktis Grup Ransomware / APT.",
         ACCENT_BLUE, Inches(3.75)),
        ("TACTICAL CTI", "Analis Forensik & Rekayasa Deteksi", "Detection Engineers, SOC L2/L3, Malware Analyst",
         "Taktik, Teknik, dan Prosedur (TTP) musuh berdasarkan framework MITRE ATT&CK. Menjawab cara musuh menembus pertahanan dan bergerak di dalam sistem.",
         "Format Output:\nYARA Rules, Aturan Deteksi Sigma, Pemetaan MITRE ATT&CK Matrix, Hunting Playbooks.",
         ACCENT_AMBER, Inches(6.7)),
        ("TECHNICAL CTI", "Garis Depan & Perimeter Keamanan", "SOC L1, Network Engineer, Firewall Admin",
         "Indikator kompromi teknis (IoC) yang berumur pendek: nilai hash berkas malware, alamat IP command-and-control, domain phishing, dan string URL jahat.",
         "Format Output:\nFeed STIX/TAXII, IP Blacklist Firewall, Hash Blocklist EDR, DNS Sinkhole entries.",
         ACCENT_GREEN, Inches(9.65))
    ]

    for title, aud_type, role_names, scope_desc, deliverable, border_c, left_p in tiers_4:
        draw_box(s4, left_p, Inches(1.6), Inches(2.8), Inches(5.3), COLOR_SURFACE, border_c)
        box = s4.shapes.add_textbox(left_p + Inches(0.15), Inches(1.75), Inches(2.5), Inches(5.0))
        tf = box.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = border_c

        p_aud = tf.add_paragraph()
        p_aud.text = aud_type.upper()
        p_aud.font.name = "Segoe UI"
        p_aud.font.size = Pt(8.5)
        p_aud.font.bold = True
        p_aud.font.color.rgb = COLOR_TEXT_MUTED
        p_aud.space_before = Pt(2)

        p_roles = tf.add_paragraph()
        p_roles.text = f"Audiens: {role_names}"
        p_roles.font.name = "Segoe UI"
        p_roles.font.size = Pt(9.5)
        p_roles.font.bold = True
        p_roles.font.color.rgb = COLOR_TEXT_PRIMARY
        p_roles.space_before = Pt(6)

        p_scope = tf.add_paragraph()
        p_scope.text = scope_desc
        p_scope.font.name = "Segoe UI"
        p_scope.font.size = Pt(9.5)
        p_scope.font.color.rgb = COLOR_TEXT_SECONDARY
        p_scope.space_before = Pt(8)

        p_deliv = tf.add_paragraph()
        p_deliv.text = deliverable
        p_deliv.font.name = "Segoe UI"
        p_deliv.font.size = Pt(9)
        p_deliv.font.color.rgb = border_c
        p_deliv.space_before = Pt(12)

    # =========================================================================
    # SLIDE 5: SIKLUS HIDUP CTI (THE 6-PHASE THREAT INTEL LIFECYCLE)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_canvas(s5)
    render_header(s5, "Metodologi Operasional", "Siklus Hidup Intelijen Ancaman: Dari Penetapan Prioritas Hingga Evaluasi Timbal-Balik")

    # 6 Steps in 2 rows of 3 columns
    cycle_steps = [
        ("01. PLANNING & DIRECTION", "Menetapkan Kebutuhan (PIR)", 
         "Mendefinisikan Priority Intelligence Requirements (PIR): Aset apa yang paling berharga bagi bisnis? Sistem mana yang paling rawan? Mencegah tim mengumpulkan data yang tidak relevan.", Inches(0.8), Inches(1.6)),
        ("02. COLLECTION", "Pengumpulan Sumber Data", 
         "Menarik telemetri dari sumber internal (sensor jaringan, honeypot, event log SIEM) dan sumber eksternal (OSINT, dark web telemetry, feed komersial, laporan ISAC).", Inches(4.75), Inches(1.6)),
        ("03. PROCESSING", "Normalisasi & Ekstraksi", 
         "Mengubah data mentah menjadi format siap analisis: de-obfuskasi skrip, ekstraksi metadata berkas, lookup rDNS, dan penghitungan checksum hash kriptografis.", Inches(8.7), Inches(1.6)),
        ("04. ANALYSIS & PRODUCTION", "Sintesis & Kontekstualisasi", 
         "Menguji hipotesis ancaman: menghubungkan anomali koneksi dengan taktik penyerang, menilai tingkat kredibilitas sumber, dan menyusun laporan berbasis bukti.", Inches(0.8), Inches(3.9)),
        ("05. DISSEMINATION", "Distribusi Tepat Sasaran", 
         "Menyalurkan hasil intelijen ke pihak yang tepat: ringkasan risiko ke pimpinan, peringatan taktis ke analis SOC L2, dan aturan IoC otomatis ke firewall/EDR.", Inches(4.75), Inches(3.9)),
        ("06. FEEDBACK & REVIEW", "Evaluasi & Kalibrasi Ulang", 
         "Menilai efektivitas tindakan: Apakah intelijen ini berhasil memblokir serangan nyata? Apakah ada celah deteksi yang terlewat? Hasil evaluasi menjadi input bagi siklus berikutnya.", Inches(8.7), Inches(3.9)),
    ]

    for step_num, step_name, step_desc, left_p, top_p in cycle_steps:
        draw_box(s5, left_p, top_p, Inches(3.8), Inches(2.1), COLOR_SURFACE, COLOR_BORDER)
        box = s5.shapes.add_textbox(left_p + Inches(0.15), top_p + Inches(0.12), Inches(3.5), Inches(1.85))
        tf = box.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = step_num
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT_CYAN

        p1 = tf.add_paragraph()
        p1.text = step_name
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_PRIMARY
        p1.space_before = Pt(2)

        p2 = tf.add_paragraph()
        p2.text = step_desc
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(9.2)
        p2.font.color.rgb = COLOR_TEXT_SECONDARY
        p2.space_before = Pt(5)

    # Analyst Field Note Banner
    draw_box(s5, Inches(0.8), Inches(6.15), Inches(11.7), Inches(0.75), COLOR_SURFACE_ALT, ACCENT_AMBER)
    fn_box = s5.shapes.add_textbox(Inches(1.0), Inches(6.2), Inches(11.3), Inches(0.65))
    tf_fn = fn_box.text_frame
    p_fn = tf_fn.paragraphs[0]
    p_fn.text = "CATATAN PRAKTISI: KESALAHAN TERBESAR PADA CTI IMPLEMENTASI"
    p_fn.font.name = "Segoe UI"
    p_fn.font.size = Pt(9)
    p_fn.font.bold = True
    p_fn.font.color.rgb = ACCENT_AMBER

    p_fnd = tf_fn.add_paragraph()
    p_fnd.text = "Lebih dari 75% organisasi langsung melompat ke tahap Koleksi (membeli feed komersial mahal) tanpa memiliki PIR yang jelas pada tahap Planning. Akibatnya, tim analis tenggelam dalam lautan data tanpa relevansi bisnis."
    p_fnd.font.name = "Segoe UI"
    p_fnd.font.size = Pt(9.5)
    p_fnd.font.color.rgb = COLOR_TEXT_PRIMARY

    # =========================================================================
    # SLIDE 6: THE DIAMOND MODEL OF INTRUSION ANALYSIS
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_canvas(s6)
    render_header(s6, "Model Analisis Ancaman", "The Diamond Model: Memetakan Hubungan Struktural dalam Analisis Intrusi")

    # Center relationship box
    draw_box(s6, Inches(4.7), Inches(2.9), Inches(3.9), Inches(1.8), COLOR_SURFACE_ALT, ACCENT_CYAN, 2)
    c_box = s6.shapes.add_textbox(Inches(4.8), Inches(3.0), Inches(3.7), Inches(1.6))
    tf_c = c_box.text_frame
    tf_c.word_wrap = True
    pc = tf_c.paragraphs[0]
    pc.text = "AKSIOMA DASAR DIAMOND MODEL"
    pc.font.name = "Segoe UI"
    pc.font.size = Pt(10)
    pc.font.bold = True
    pc.font.color.rgb = ACCENT_CYAN
    pc.alignment = PP_ALIGN.CENTER

    pc2 = tf_c.add_paragraph()
    pc2.text = "\"Setiap peristiwa intrusi ditandai oleh Pelaku (Adversary) yang mengerahkan Senjata/Alat (Capability) melalui Saluran (Infrastructure) terhadap sebuah Sasaran (Victim).\""
    pc2.font.name = "Segoe UI"
    pc2.font.size = Pt(9.5)
    pc2.font.color.rgb = COLOR_TEXT_PRIMARY
    pc2.alignment = PP_ALIGN.CENTER
    pc2.space_before = Pt(4)

    # 4 Nodes of the Diamond
    nodes_diamond = [
        ("ADVERSARY (PELAKU)", "Identitas atau organisasi di balik serangan.",
         "• Meneliti motivasi: spionase, finansial, reputasi\n• Mengidentifikasi operator versus sponsor pendana\n• Melacak kebiasaan jam operasional penyerang",
         Inches(4.7), Inches(1.5), Inches(3.9), Inches(1.3), ACCENT_RED),
        ("CAPABILITY (KAPABILITAS / ALAT)", "Senjata dan teknik perangkat lunak yang digunakan.",
         "• Exploit zero-day atau 1-day pada celah software\n• Custom implant, webshell, atau skrip otomatisasi\n• Efektivitas teknik evasif terhadap antivirus/EDR",
         Inches(0.8), Inches(3.1), Inches(3.7), Inches(1.4), ACCENT_AMBER),
        ("INFRASTRUCTURE (INFRASTRUKTUR)", "Jalur komunikasi dan server yang dikendalikan.",
         "• Alamat IP, server proxy, relay VPN, bulletproof host\n• Server Command-and-Control (C2), domain registrar\n• Pemanfaatan botnet perantara untuk kamuflase",
         Inches(8.8), Inches(3.1), Inches(3.7), Inches(1.4), ACCENT_BLUE),
        ("VICTIM (SASARAN / KORBAN)", "Target organisasi atau aset yang disasar.",
         "• Personel kunci: staf keuangan, admin sistem IT\n• Infrastruktur kritis: server basis data, API gateway\n• Nilai strategis aset bagi kelangsungan bisnis",
         Inches(4.7), Inches(4.85), Inches(3.9), Inches(1.3), ACCENT_GREEN),
    ]

    for title, role_s, bullets, left_p, top_p, width_p, height_p, col_c in nodes_diamond:
        draw_box(s6, left_p, top_p, width_p, height_p, COLOR_SURFACE, col_c)
        box = s6.shapes.add_textbox(left_p + Inches(0.12), top_p + Inches(0.08), width_p - Inches(0.24), height_p - Inches(0.16))
        tf = box.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = col_c

        p1 = tf.add_paragraph()
        p1.text = role_s
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(9)
        p1.font.color.rgb = COLOR_TEXT_PRIMARY
        p1.space_before = Pt(1)

        p2 = tf.add_paragraph()
        p2.text = bullets
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = COLOR_TEXT_SECONDARY
        p2.space_before = Pt(3)

    # Contextual axes footer note
    foot_box = s6.shapes.add_textbox(Inches(0.8), Inches(6.3), Inches(11.733), Inches(0.5))
    tf_ft = foot_box.text_frame
    p = tf_ft.paragraphs[0]
    p.text = "DUA SUMBU PENDUKUNG MODEL DIAMOND: Sumbu Sosial-Politik (menghubungkan Pelaku dengan Sasaran) dan Sumbu Teknologi (menghubungkan Kapabilitas dengan Infrastruktur)."
    p.font.name = "Segoe UI"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 7: CYBER KILL CHAIN & DEFENSIVE LEVERAGE
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    apply_canvas(s7)
    render_header(s7, "Model Analisis Ancaman", "Cyber Kill Chain: Anatomi Intrusi dan Titik Intervensi Defensif")

    # 7 Progressive phases horizontally
    phases_kc = [
        ("FASE 01", "Reconnaissance", "Pengintaian", "Pemindaian port, enumerasi email target, pengumpulan OSINT.", ACCENT_CYAN, "Leverage: TERTINGGI\nBiaya Mitigasi: Minimal"),
        ("FASE 02", "Weaponization", "Persenjataan", "Penyisipan muatan exploit ke dalam berkas dokumen atau installer.", ACCENT_CYAN, "Analisis di sandbox / feed eksternal"),
        ("FASE 03", "Delivery", "Pengiriman", "Penyebaran melalui spear-phishing email atau layanan web terbuka.", ACCENT_BLUE, "Email gateway & firewall filtering"),
        ("FASE 04", "Exploitation", "Eksploitasi", "Celah software dieksekusi saat berkas dibuka oleh korban.", ACCENT_AMBER, "Patch management & EDR block"),
        ("FASE 05", "Installation", "Instalasi", "Pemasangan backdoor atau mekanisme persistensi di sistem operasi.", ACCENT_AMBER, "Integritas file & hak akses admin"),
        ("FASE 06", "C2 Channel", "Komunikasi C2", "Malware menghubungi server pengendali milik penyerang.", ACCENT_RED, "DNS sinkhole & proxy egress filter"),
        ("FASE 07", "Actions on Obj.", "Tujuan Akhir", "Pencurian basis data sensitif, sabotase, atau ransomware.", ACCENT_RED, "Leverage: TERENDAH\nDampak Kerugian: Fatal")
    ]

    for i, (f_num, f_en, f_id, f_desc, f_col, f_lev) in enumerate(phases_kc):
        left_pos = Inches(0.8 + i * 1.68)
        draw_box(s7, left_pos, Inches(1.6), Inches(1.58), Inches(4.5), COLOR_SURFACE, f_col)
        box = s7.shapes.add_textbox(left_pos + Inches(0.08), Inches(1.7), Inches(1.42), Inches(4.3))
        tf = box.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = f_num
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(9)
        p0.font.bold = True
        p0.font.color.rgb = f_col

        p1 = tf.add_paragraph()
        p1.text = f_en
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(10.5)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_PRIMARY
        p1.space_before = Pt(3)

        p1_sub = tf.add_paragraph()
        p1_sub.text = f"({f_id})"
        p1_sub.font.name = "Segoe UI"
        p1_sub.font.size = Pt(8.5)
        p1_sub.font.color.rgb = COLOR_TEXT_MUTED
        p1_sub.space_before = Pt(1)

        p2 = tf.add_paragraph()
        p2.text = f_desc
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(8.8)
        p2.font.color.rgb = COLOR_TEXT_SECONDARY
        p2.space_before = Pt(8)

        p3 = tf.add_paragraph()
        p3.text = f_lev
        p3.font.name = "Segoe UI"
        p3.font.size = Pt(8)
        p3.font.bold = True
        p3.font.color.rgb = f_col
        p3.space_before = Pt(12)

    # Bottom takeaway
    draw_box(s7, Inches(0.8), Inches(6.25), Inches(11.7), Inches(0.65), COLOR_SURFACE_ALT, COLOR_BORDER)
    tb_bot = s7.shapes.add_textbox(Inches(1.0), Inches(6.28), Inches(11.3), Inches(0.6))
    tf_b = tb_bot.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "PRINSIP ASIMETRI DEFENSIP: Menggagalkan penyerang pada 1 tahap mana pun dalam Kill Chain akan membatalkan seluruh operasi mereka. Namun, jika penyerang mencapai Fase 07, kerugian bisnis organisasi telah terjadi."
    p_b.font.name = "Segoe UI"
    p_b.font.size = Pt(9.2)
    p_b.font.color.rgb = COLOR_TEXT_PRIMARY

    # =========================================================================
    # SLIDE 8: THE PYRAMID OF PAIN (DAVID BIANCO)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_canvas(s8)
    render_header(s8, "Prinsip Nilai Intelijen", "Pyramid of Pain: Mengukur Dampak dan Biaya Hambatan bagi Penyerang")

    # 6 Layers from Top (Tough) to Bottom (Trivial)
    pop_layers = [
        ("TTPs (Taktik, Teknik, dan Prosedur)", "BIAYA BAGI PENYERANG: SANGAT TINGGI (TOUGH)",
         "Memaksa penyerang merombak total metodologi, melatih ulang personel peretas, atau mencari vektor serangan baru.",
         ACCENT_RED, Inches(1.6)),
        ("Tools / Perangkat Lunak Serangan", "BIAYA BAGI PENYERANG: TINGGI (CHALLENGING)",
         "Memblokir utility atau builder malware memaksa lawan merekayasa ulang software exploit atau membeli tool baru.",
         ACCENT_AMBER, Inches(2.45)),
        ("Network / Host Artifacts", "BIAYA BAGI PENYERANG: MENENGAH (ANNOYING)",
         "Mendeteksi pola URI khas, metadata kompilasi (PDB path), nama author dokumen decoy, atau string User-Agent spesifik.",
         ACCENT_PURPLE, Inches(3.3)),
        ("Domain Names", "BIAYA BAGI PENYERANG: RENDAH (SIMPLE)",
         "Penyerang dapat dengan mudah mendaftarkan nama domain baru melalui registrar yang longgar dengan biaya murah.",
         ACCENT_BLUE, Inches(4.15)),
        ("Alamat IP (IP Addresses)", "BIAYA BAGI PENYERANG: SANGAT RENDAH (EASY)",
         "Penyerang dapat dengan mudah berganti alamat IP publik melalui proxy komersial, VPS sekali pakai, atau node VPN.",
         ACCENT_CYAN, Inches(5.0)),
        ("Nilai Hash Kriptografis (MD5 / SHA-256)", "BIAYA BAGI PENYERANG: TIDAK BERPENGARUH (TRIVIAL)",
         "Menambahkan 1 karakter kosong (whitespace) pada kode malware akan mengubah seluruh nilai checksum hash berkas.",
         COLOR_TEXT_MUTED, Inches(5.85))
    ]

    for label_name, impact_cost, expl, col_layer, top_pos in pop_layers:
        draw_box(s8, Inches(0.8), top_pos, Inches(11.733), Inches(0.72), COLOR_SURFACE, col_layer)
        box = s8.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.04), Inches(11.3), Inches(0.65))
        tf = box.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = f"{label_name}   |   {impact_cost}"
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(10.5)
        p0.font.bold = True
        p0.font.color.rgb = col_layer

        p1 = tf.add_paragraph()
        p1.text = expl
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(9.2)
        p1.font.color.rgb = COLOR_TEXT_PRIMARY
        p1.space_before = Pt(1)

    # Footnote
    fn_pop = s8.shapes.add_textbox(Inches(0.8), Inches(6.65), Inches(11.7), Inches(0.35))
    tf_fpop = fn_pop.text_frame
    p_fpop = tf_fpop.paragraphs[0]
    p_fpop.text = "KESIMPULAN: Kebanyakan SOC menghabiskan 90% waktu memblokir Hash & IP (Trivial/Easy), padahal nilai pertahanan tertinggi berada pada TTPs."
    p_fpop.font.name = "Segoe UI"
    p_fpop.font.size = Pt(8.5)
    p_fpop.font.bold = True
    p_fpop.font.color.rgb = ACCENT_AMBER

    # =========================================================================
    # SLIDE 9: ACTIVE DEFENSE & DECEPTION ARCHITECTURE
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    apply_canvas(s9)
    render_header(s9, "Teknologi Intelijen Modern", "Active Defense: Mengubah Jebakan Deception Menjadi Sumber Intelijen Berkualitas Tinggi")

    # Left Box: Konsep Deception
    draw_box(s9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3), COLOR_SURFACE, COLOR_BORDER)
    tb_dec = s9.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.9))
    tf_dec = tb_dec.text_frame
    tf_dec.word_wrap = True

    p = tf_dec.paragraphs[0]
    p.text = "ARSITEKTUR DECEPTION & TRIPWIRE"
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    deception_points = [
        ("Penyebaran Honeytoken", "Menempatkan kredensial palsu (misal API key atau akun service palsu) dan dokumen umpan pada titik strategis yang tidak memiliki aktivitas operasional sah."),
        ("Tingkat Kepastian 100% (Zero False Positive)", "Pengguna legal dan sistem produksi tidak memiliki alasan untuk mengakses aset jebakan. Akses apa pun langsung membuktikan adanya anomali atau intrusi."),
        ("Mengukur Waktu Dwell Time Musuh", "Memberikan visibilitas instan terhadap kehadiran penyerang sebelum mereka berhasil menemukan dan mengeksfiltrasi data produksi riil.")
    ]
    for tit, dsc in deception_points:
        pt = tf_dec.add_paragraph()
        pt.text = tit
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(11.5)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_TEXT_PRIMARY
        pt.space_before = Pt(14)

        pd = tf_dec.add_paragraph()
        pd.text = dsc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(10)
        pd.font.color.rgb = COLOR_TEXT_SECONDARY
        pd.space_before = Pt(3)

    # Right Box: Nilai CTI dari Deception
    draw_box(s9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3), COLOR_SURFACE, COLOR_BORDER)
    tb_val = s9.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.9))
    tf_val = tb_val.text_frame
    tf_val.word_wrap = True

    p = tf_val.paragraphs[0]
    p.text = "NILAI INTELIJEN LANGSUNG BAGI SOC"
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    value_points = [
        ("Observasi Taktik Langsung (In-Situ Observation)", "Mengamati bagaimana penyerang berinteraksi dengan server umpan: perintah terminal apa yang diketik, alat apa yang diunduh, dan apa target berikutnya."),
        ("Membalikkan Keuntungan Asimetri", "Penyerang terpaksa menghabiskan waktu dan sumber daya pada sistem decoy, sementara tim bertahan mendapatkan waktu untuk isolasi dan analisis."),
        ("Penghasil Indikator Kredibel (High-Fidelity IOC)", "Semua hash, IP, dan domain yang diekstraksi dari honeypot terkonfirmasi jahat tanpa keraguan, memungkinkan blokir otomatis secara instan.")
    ]
    for tit, dsc in value_points:
        pt = tf_val.add_paragraph()
        pt.text = tit
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(11.5)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_TEXT_PRIMARY
        pt.space_before = Pt(14)

        pd = tf_val.add_paragraph()
        pd.text = dsc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(10)
        pd.font.color.rgb = COLOR_TEXT_SECONDARY
        pd.space_before = Pt(3)

    # =========================================================================
    # SLIDE 10: DE-ANONYMIZATION & FORENSIC ATTRIBUTION
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    apply_canvas(s10)
    render_header(s10, "Investigasi Forensik", "Metodologi Atribusi & De-Anonymization: Membongkar Kamuflase Penyerang")

    # 4 Attribution Pillars
    pillars = [
        ("01. KEBOCORAN PROTOKOL WEBRTC", "Membongkar IP Asli di Balik VPN",
         "Penyerang sering menggunakan proxy komersial atau VPN. Melalui sesi WebRTC yang sah, browser dapat mengekstrak host candidate yang mengungkap IP internal LAN penyerang atau alamat publik sekunder yang tidak terproteksi VPN.",
         ACCENT_CYAN, Inches(0.8), Inches(1.6)),
        ("02. SIDIK JARI LAPISAN JARINGAN (L4)", "Fingerprinting Stack TCP / SYN",
         "Karakteristik paket SYN awal (kombinasi TTL, Window Size, urutan TCP Options, dan ukuran MSS) membentuk tanda khas sistem operasi atau alat scanning khusus (misalnya Nmap, Masscan, atau Go-custom scanner).",
         ACCENT_BLUE, Inches(6.8), Inches(1.6)),
        ("03. METADATA & ARTEFAK DOKUMEN", "Jejak Lingkungan Pengembangan",
         "Analisis forensik terhadap berkas yang ditinggalkan: author name dokumen Office, jalur kompilasi debug (PDB path), locale bahasa keyboard, serta timestamp zona waktu sistem penyerang.",
         ACCENT_AMBER, Inches(0.8), Inches(4.1)),
        ("04. SINKRONISASI CANARY BEACON", "Mengeksploitasi Kesalahan Operasional (OPSEC)",
         "Ketika penyerang mengunduh dokumen canary lewat VPN, lalu memindahkannya dan membukanya pada komputer kerja lokal tanpa VPN aktif, beacon rahasia akan menembak ke server atribusi dan membongkar identitas jaringan aslinya.",
         ACCENT_GREEN, Inches(6.8), Inches(4.1))
    ]

    for p_num, p_sub, p_body, col_p, left_p, top_p in pillars:
        draw_box(s10, left_p, top_p, Inches(5.7), Inches(2.3), COLOR_SURFACE, col_p)
        box = s10.shapes.add_textbox(left_p + Inches(0.15), top_p + Inches(0.12), Inches(5.4), Inches(2.05))
        tf = box.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = p_num
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = col_p

        p1 = tf.add_paragraph()
        p1.text = p_sub
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_PRIMARY
        p1.space_before = Pt(2)

        p2 = tf.add_paragraph()
        p2.text = p_body
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(9.2)
        p2.font.color.rgb = COLOR_TEXT_SECONDARY
        p2.space_before = Pt(6)

    # Bottom Reality Check
    draw_box(s10, Inches(0.8), Inches(6.55), Inches(11.7), Inches(0.45), COLOR_SURFACE_ALT, COLOR_BORDER)
    foot_at = s10.shapes.add_textbox(Inches(1.0), Inches(6.58), Inches(11.3), Inches(0.4))
    tf_fat = foot_at.text_frame
    pf = tf_fat.paragraphs[0]
    pf.text = "PRINSIP FORENSIK: Atribusi bukanlah tebakan tunggal, melainkan akumulasi korelasi teknis dari berbagai lapisan (L4 hingga L7) yang saling menguatkan bukti hukum."
    pf.font.name = "Segoe UI"
    pf.font.size = Pt(8.5)
    pf.font.bold = True
    pf.font.color.rgb = COLOR_TEXT_MUTED

    # =========================================================================
    # SLIDE 11: STANDARDISASI & PERTUKARAN CTI (STIX 2.1 & TAXII)
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    apply_canvas(s11)
    render_header(s11, "Interoperabilitas Industri", "Standardisasi Berbagi Intelijen: Kerangka STIX 2.1 dan Protokol Transportasi TAXII")

    # Left: STIX 2.1
    draw_box(s11, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3), COLOR_SURFACE, COLOR_BORDER)
    tb_st = s11.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.9))
    tf_st = tb_st.text_frame
    tf_st.word_wrap = True

    p = tf_st.paragraphs[0]
    p.text = "STIX 2.1 (STRUKTUR DATA INTELIJEN)"
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    stix_bullets = [
        ("Format Graf Berbasis Objek (JSON)", "Structured Threat Information Expression (STIX) mendefinisikan bahasa standar yang merepresentasikan relasi ancaman dalam bentuk graf."),
        ("Entitas Kunci (SDO / SRO):", "• Threat-Actor: Siapa dalang penyerangan\n• Attack-Pattern: Taktik & teknik MITRE ATT&CK\n• Indicator: Alamat IP, hash, pola regex pendeteksi\n• Relationship: Menghubungkan aktor -> teknik -> indikator"),
        ("Manfaat Bagi Ekosistem:", "Menghilangkan ambiguitas laporan manual berbasis PDF dan memungkinkan integrasi data antar produk keamanan yang berbeda vendor.")
    ]
    for tit, dsc in stix_bullets:
        pt = tf_st.add_paragraph()
        pt.text = tit
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(11)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_TEXT_PRIMARY
        pt.space_before = Pt(14)

        pd = tf_st.add_paragraph()
        pd.text = dsc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = COLOR_TEXT_SECONDARY
        pd.space_before = Pt(3)

    # Right: TAXII 2.1
    draw_box(s11, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3), COLOR_SURFACE, COLOR_BORDER)
    tb_tx = s11.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.9))
    tf_tx = tb_tx.text_frame
    tf_tx.word_wrap = True

    p = tf_tx.paragraphs[0]
    p.text = "TAXII 2.1 (PROTOKOL PENGIRIMAN AMAN)"
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_PURPLE

    taxii_bullets = [
        ("Transport Berbasis HTTPS REST API", "Trusted Automated Exchange of Intelligence Information (TAXII) adalah protokol pengangkutan pesan STIX secara terenkripsi."),
        ("Dua Model Pertukaran:", "• Channel Model: Siaran pub/sub satu arah dari penyedia feed ke pelanggan\n• Collection Model: Repositori dua arah di mana anggota komunitas saling mengunggah dan menarik berkas intelijen"),
        ("Konektivitas Kolaboratif Sektor Finansial & BSSN:", "Memungkinkan bank, badan siber pemerintah, dan industri saling bertukar ancaman terbaru dalam waktu sub-detik secara otomatis tanpa intervensi manual.")
    ]
    for tit, dsc in taxii_bullets:
        pt = tf_tx.add_paragraph()
        pt.text = tit
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(11)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_TEXT_PRIMARY
        pt.space_before = Pt(14)

        pd = tf_tx.add_paragraph()
        pd.text = dsc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = COLOR_TEXT_SECONDARY
        pd.space_before = Pt(3)

    # =========================================================================
    # SLIDE 12: OTOMASI & ANALISIS MODERN (HUMAN-IN-THE-LOOP AGENTIC CTI)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    apply_canvas(s12)
    render_header(s12, "Evolusi Operasional", "Otomasi & Multi-Agent CTI: Meningkatkan Kecepatan Triage Tanpa Menghilangkan Kontrol Analis")

    # 3 Workflow Comparison Columns
    workflows = [
        ("PROSES TRADISIONAL (MANUAL)", "Waktu Triage: 2 - 4 Jam",
         "1. Analis menerima alert mentah dari SIEM.\n2. Melakukan query manual ke VirusTotal, AbuseIPDB, dan Shodan.\n3. Membuka laporan MITRE secara manual.\n4. Mengetik draf ringkasan ke tiket insiden.\n\nKelemahan: Bottleneck manusia, keterlambatan mitigasi, dan kelelahan mental tim SOC.",
         ACCENT_RED, Inches(0.8)),
        ("ALUR OTOMASI MULTI-AGENT", "Waktu Triage: < 30 Detik",
         "1. Data telemetry masuk ke bus in-memory.\n2. Agen Recon: Mengekstrak kecepatan scan L4 & rDNS.\n3. Agen DFIR: Mengekstrak metadata berkas & checksum.\n4. Agen Atribusi: Memvalidasi korelasi token & IP.\n5. Agen Strategi: Menghubungkan teknik ke MITRE.\n\nKeunggulan: Korelasi instan dengan data historis.",
         ACCENT_CYAN, Inches(4.75)),
        ("HUMAN-IN-THE-LOOP CONTROL", "Keputusan Akhir di Tangan Analis",
         "1. Sistem menyajikan draf investigasi terstruktur lengkap dengan bukti teknis.\n2. Analis senior memvalidasi rekomendasi tindakan.\n3. Eksekusi firewall rule dan isolasi perangkat dilakukan dengan satu klik konfirmasi.\n\nHasil: Produktivitas naik 10x dengan risiko mitigasi salah (false-block) terkendali.",
         ACCENT_GREEN, Inches(8.7))
    ]

    for title, metric, body_desc, col_c, left_p in workflows:
        draw_box(s12, left_p, Inches(1.6), Inches(3.8), Inches(4.2), COLOR_SURFACE, col_c)
        box = s12.shapes.add_textbox(left_p + Inches(0.15), Inches(1.75), Inches(3.5), Inches(3.9))
        tf = box.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = col_c

        p_m = tf.add_paragraph()
        p_m.text = metric
        p_m.font.name = "Segoe UI"
        p_m.font.size = Pt(9.5)
        p_m.font.bold = True
        p_m.font.color.rgb = COLOR_TEXT_PRIMARY
        p_m.space_before = Pt(3)

        p_b = tf.add_paragraph()
        p_b.text = body_desc
        p_b.font.name = "Segoe UI"
        p_b.font.size = Pt(9.2)
        p_b.font.color.rgb = COLOR_TEXT_SECONDARY
        p_b.space_before = Pt(8)

    # Executive Closing Callout
    draw_box(s12, Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.9), COLOR_SURFACE_ALT, ACCENT_AMBER)
    box_c = s12.shapes.add_textbox(Inches(1.0), Inches(6.08), Inches(11.3), Inches(0.75))
    tf_c = box_c.text_frame
    p0 = tf_c.paragraphs[0]
    p0.text = "RINGKASAN EKSEKUTIF CTI"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(9)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_AMBER

    p1 = tf_c.add_paragraph()
    p1.text = "Cyber Threat Intelligence yang efektif bukan tentang seberapa banyak data yang dikumpulkan, melainkan seberapa cepat intelijen tersebut membantu analis mengambil keputusan yang tepat untuk menghentikan musuh sebelum dampak finansial terjadi."
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(9.8)
    p1.font.color.rgb = COLOR_TEXT_PRIMARY
    p1.space_before = Pt(2)

    output_path = "CTI_Fundamentals_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation successfully written to {output_path}")

if __name__ == "__main__":
    build_human_cti_deck()
